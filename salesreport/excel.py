"""Writing the formatted Excel report."""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4D3A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=18)
SECTION_FONT = Font(bold=True, size=13)
MUTED_FONT = Font(color="666666")

KR = '#,##0 "kr"'
KR_DEC = '#,##0.00 "kr"'
INT = "#,##0"
PCT = "0.0%"
DATE = "DD.MM.YYYY"


def _write_table(ws, top, left, headers, rows, formats=None):
    """Write a header row plus data rows. Returns the last row number used."""
    formats = formats or {}
    for j, title in enumerate(headers):
        cell = ws.cell(row=top, column=left + j, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for i, values in enumerate(rows, start=1):
        for j, value in enumerate(values):
            cell = ws.cell(row=top + i, column=left + j, value=value)
            if headers[j] in formats:
                cell.number_format = formats[headers[j]]
    return top + len(rows)


def _fit_columns(ws, min_width=10, max_width=48):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = cell.value.strftime("%d.%m.%Y") if isinstance(cell.value, datetime) else str(cell.value)
            if isinstance(cell.value, float):
                text = f"{cell.value:,.0f} kr"
            widths[cell.column] = max(widths.get(cell.column, 0), len(text) + 2)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(width, max_width))


def _breakdown(clean, column):
    total = clean["revenue"].sum()
    grouped = (
        clean.groupby(column)
        .agg(orders=("order_id", "nunique"), units=("quantity", "sum"), revenue=("revenue", "sum"))
        .sort_values("revenue", ascending=False)
    )
    return [
        [name, int(r.orders), int(r.units), float(r.revenue), float(r.revenue / total) if total else 0.0]
        for name, r in grouped.iterrows()
    ]


def _data_sheet(wb, title, headers, rows, formats):
    ws = wb.create_sheet(title)
    last = _write_table(ws, 1, 1, headers, rows, formats)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(last, 2)}"
    _fit_columns(ws)
    return ws


def write_report(clean, issues, files, path, title="Sales report"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    revenue = float(clean["revenue"].sum())
    orders = int(clean["order_id"].nunique())
    units = int(clean["quantity"].sum())
    first, last = clean["date"].min(), clean["date"].max()

    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Period {first:%d.%m.%Y} to {last:%d.%m.%Y}. Built from {len(files)} files on {datetime.now():%d.%m.%Y %H:%M}."
    ws["A2"].font = MUTED_FONT

    kpis = [
        ("Revenue", revenue, KR),
        ("Orders", orders, INT),
        ("Units sold", units, INT),
        ("Average order value", revenue / orders if orders else 0, KR),
        ("Rows used", len(clean), INT),
        ("Rows set aside (see Issues)", len(issues), INT),
    ]
    for i, (label, value, fmt) in enumerate(kpis, start=4):
        ws.cell(row=i, column=1, value=label)
        cell = ws.cell(row=i, column=2, value=value)
        cell.number_format = fmt
        cell.font = Font(bold=True)

    # Revenue by month, with a chart next to it.
    ws.cell(row=11, column=1, value="Revenue by month").font = SECTION_FONT
    monthly = (
        clean.groupby("month")
        .agg(orders=("order_id", "nunique"), units=("quantity", "sum"), revenue=("revenue", "sum"))
        .sort_index()
    )
    month_rows = [
        [datetime.strptime(m, "%Y-%m").strftime("%b %Y"), int(r.orders), int(r.units), float(r.revenue)]
        for m, r in monthly.iterrows()
    ]
    month_end = _write_table(ws, 12, 1, ["Month", "Orders", "Units", "Revenue"], month_rows, {"Revenue": KR, "Units": INT})

    chart = BarChart()
    chart.title = "Revenue by month"
    chart.y_axis.title = "kr"
    chart.y_axis.numFmt = "#,##0"
    chart.legend = None
    chart.height, chart.width = 7.5, 15
    chart.add_data(Reference(ws, min_col=4, min_row=12, max_row=month_end), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=13, max_row=month_end))
    chart.series[0].graphicalProperties.solidFill = "1F4D3A"
    ws.add_chart(chart, "F4")

    # Top products
    top_start = month_end + 3
    ws.cell(row=top_start - 1, column=1, value="Top 5 products").font = SECTION_FONT
    top = _breakdown(clean, "product")[:5]
    _write_table(ws, top_start, 1, ["Product", "Orders", "Units", "Revenue", "Share"], top, {"Revenue": KR, "Share": PCT})
    _fit_columns(ws)
    ws.column_dimensions["A"].width = 30

    breakdown_headers = ["Name", "Orders", "Units", "Revenue", "Share of revenue"]
    breakdown_formats = {"Revenue": KR, "Share of revenue": PCT, "Units": INT}
    for sheet, column in (("By product", "product"), ("By city", "city"), ("By channel", "channel")):
        headers = [sheet.split()[-1].capitalize()] + breakdown_headers[1:]
        formats = dict(breakdown_formats)
        _data_sheet(wb, sheet, headers, _breakdown(clean, column), formats)

    data_rows = [
        [datetime.combine(r.date, datetime.min.time()), r.order_id, r.city, r.product, int(r.quantity),
         float(r.unit_price), float(r.revenue), r.channel, r.source_file, int(r.source_row)]
        for r in clean.itertuples()
    ]
    _data_sheet(
        wb, "Data",
        ["Date", "Order", "City", "Product", "Quantity", "Unit price", "Revenue", "Channel", "Source file", "Source row"],
        data_rows,
        {"Date": DATE, "Unit price": KR_DEC, "Revenue": KR_DEC},
    )

    issue_rows = [[r.source_file, int(r.source_row), r.reason, r.order_id, r.product] for r in issues.itertuples()]
    _data_sheet(wb, "Issues", ["Source file", "Source row", "Reason", "Order", "Product"], issue_rows, {})

    wb.save(path)
