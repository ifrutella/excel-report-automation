from datetime import date, datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from salesreport.clean import clean, load_folder, parse_date, parse_number, product_key, standardize
from salesreport.cli import main

SAMPLE = Path(__file__).resolve().parent.parent / "data" / "raw"


@pytest.mark.parametrize(
    "text, expected",
    [
        ("1 234,50", 1234.5),
        ("1 234,50", 1234.5),
        ("1,234.50", 1234.5),
        ("1.234,50", 1234.5),
        ("169.00", 169.0),
        ("169,00", 169.0),
        ("429 kr", 429.0),
        ("NOK 59", 59.0),
        ("1.234.567", 1234567.0),
        (12, 12.0),
        ("", None),
        ("n/a", None),
        (None, None),
        (float("nan"), None),
    ],
)
def test_parse_number(text, expected):
    assert parse_number(text) == expected


@pytest.mark.parametrize(
    "value, expected",
    [
        ("2026-07-14", date(2026, 7, 14)),
        ("14.07.2026", date(2026, 7, 14)),
        ("14/07/2026", date(2026, 7, 14)),
        (datetime(2026, 7, 14, 9, 30), date(2026, 7, 14)),
        (pd.Timestamp("2026-07-14"), date(2026, 7, 14)),
        ("31.02.2026", None),
        ("soon", None),
        ("", None),
    ],
)
def test_parse_date(value, expected):
    assert parse_date(value) == expected


def test_product_key_ignores_case_spacing_and_unit_gap():
    assert product_key("  ETHIOPIA  Yirgacheffe 250 g") == product_key("Ethiopia Yirgacheffe 250g")
    assert product_key("House espresso 1 kg") == "house espresso 1kg"


def test_standardize_maps_norwegian_headers():
    raw = pd.DataFrame([["HL-1", "01.08.2026", "Oslo", "Kaffe", "2", "149,00", "Café"]],
                       columns=["Ordrenr", "Dato", "By", "Produkt", "Antall", "Pris", "Kanal"])
    df = standardize(raw, "test.csv")
    assert list(df.columns[:7]) == ["order_id", "date", "city", "product", "quantity", "unit_price", "channel"]
    assert df.loc[0, "source_row"] == 2


def test_standardize_reports_missing_columns():
    raw = pd.DataFrame([["HL-1", "Oslo"]], columns=["Order ID", "City"])
    with pytest.raises(ValueError, match="could not find column"):
        standardize(raw, "broken.csv")


def test_clean_sets_aside_bad_rows_and_duplicates():
    raw = pd.DataFrame(
        [
            ["A1", "2026-07-01", "Oslo", "Kenya 250g", "2", "100", "Webshop"],
            ["A1", "2026-07-01", "Oslo", "kenya 250 g", "2", "100", "Webshop"],  # duplicate, other spelling
            ["A2", "not a date", "Oslo", "Kenya 250g", "1", "100", "Webshop"],
            ["A3", "2026-07-02", "Oslo", "Kenya 250g", "", "100", "Webshop"],
            ["A4", "2026-07-03", "Bergen", "Kenya 250g", "3", "100,50", "Café"],
        ],
        columns=["order_id", "date", "city", "product", "quantity", "unit_price", "channel"],
    )
    data, issues = clean([standardize(raw, "t.csv")])

    assert list(data["order_id"]) == ["A1", "A4"]
    assert data["revenue"].sum() == pytest.approx(2 * 100 + 3 * 100.5)
    assert set(data["product"]) == {"Kenya 250g"}
    assert list(issues["source_row"]) == [3, 4, 5]
    assert issues.loc[0, "reason"] == "duplicate of t.csv row 2"


def test_sample_data_end_to_end(tmp_path):
    frames, files = load_folder(SAMPLE)
    data, issues = clean(frames)
    assert len(files) == 3
    assert data["product"].nunique() == 6
    assert len(data) + len(issues) == sum(len(f) for f in frames)

    out = tmp_path / "report.xlsx"
    assert main([str(SAMPLE), "-o", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "By product", "By city", "By channel", "Data", "Issues"]
    assert wb["Summary"]["B4"].value == pytest.approx(data["revenue"].sum())


def test_missing_folder_content_gives_clear_error(tmp_path, capsys):
    assert main([str(tmp_path)]) == 1
    assert "No .csv or .xlsx files" in capsys.readouterr().err
